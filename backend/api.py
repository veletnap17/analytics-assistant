import uuid
from fastapi import FastAPI
from fastapi.middleware.cors import CORSMiddleware
from fastapi.encoders import jsonable_encoder
from pydantic import BaseModel
from src.sql_generator import generate_sql
from src.services.database_service import execute_query
from src.services.result_explainer import explain_result
from src.services.source_router import detect_source
from src.services.chat_history_service import save_message, load_session, get_last_sql
from src.agents.fleet_agent import handle_fleet_question
from src.services.kpi_service import get_kpis
from src.services.chat_history_service import get_sessions, load_session
from io import BytesIO
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from src.services.followup_service import generate_followups



app = FastAPI(title="Anytime Analytics API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

class ChatRequest(BaseModel):
    question: str
    session_id: str | None = None

class ExportRequest(BaseModel):
    data: list[dict]
    filename: str = "anytime-analytics.xlsx"

@app.get("/sessions")
def sessions():
    rows = get_sessions()
    return jsonable_encoder([
        {"session_id": session_id, "title": title, "created_at": created_at}
        for session_id, title, created_at in rows
    ])

@app.get("/sessions/{session_id}")
def session(session_id: str):
    return jsonable_encoder({
        "session_id": session_id,
        "messages": load_session(session_id)
    })

@app.get("/health")
def health():
    return {"status": "ok"}

@app.get("/kpis")
def kpis():
    return get_kpis()

@app.post("/chat")
def chat(request: ChatRequest):
    session_id = request.session_id or str(uuid.uuid4())
    history = [
        {"role": m["role"], "content": m["content"]}
        for m in load_session(session_id)
    ]
    previous_sql = get_last_sql(session_id)

    save_message(session_id, "user", request.question)
    source = detect_source(request.question)

    if source == "fleet":
        answer, data = handle_fleet_question(request.question)
        suggestions = generate_followups(request.question, answer)
        if isinstance(data, dict):
            data = [{"model": k, "vehicles": v} for k, v in data.items()]

        save_message(session_id, "assistant", answer, data_json=data)

        return jsonable_encoder({
            "session_id": session_id,
            "source": "fleet",
            "answer": answer,
            "data": data,
            "sql": None,
            "suggestions": suggestions
        })

    sql = generate_sql(request.question, history, previous_sql)

    if sql == "SQL validation failed.":
        return {
            "session_id": session_id,
            "source": "sql",
            "answer": "I couldn't safely generate a query for this request.",
            "data": None,
            "sql": None
        }

    columns, rows = execute_query(sql)
    answer = explain_result(request.question, columns, rows, sql)
    suggestions = generate_followups(request.question, answer)
    data = jsonable_encoder([dict(zip(columns, row)) for row in rows])

    save_message(
        session_id,
        "assistant",
        answer,
        sql_text=sql,
        data_json=data
    )

    return jsonable_encoder({
        "session_id": session_id,
        "source": "sql",
        "answer": answer,
        "data": data,
        "sql": sql,
        "suggestions": suggestions
    })

@app.post("/export/excel")
def export_excel(request: ExportRequest):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"

    if request.data:
        columns = list(request.data[0].keys())
        ws.append(columns)

        for row in request.data:
            ws.append([row.get(col) for col in columns])

        for i, col in enumerate(columns, 1):
            width = max(len(str(col)), *(len(str(row.get(col, ""))) for row in request.data))
            ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 40)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{request.filename}"'}
    )